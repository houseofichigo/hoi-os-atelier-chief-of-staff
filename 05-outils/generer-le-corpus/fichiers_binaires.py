"""Basira Conseil demo corpus — binary files (docx, xlsx, pdf, jpg, png, zip) + client briefs.
All data fictional. Demo 'today' = Thursday 8 October 2026."""
import io
import os
import random
import zipfile
from pathlib import Path

import docx
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

OUT = Path(os.environ.get("OUT", Path(__file__).resolve().parent / "sortie" / "fichiers-en-vrac"))
OUT.mkdir(parents=True, exist_ok=True)
FICTIF = "Données fictives — démonstration Basira Conseil."
FD = os.environ.get("FONTS", str(Path(__file__).resolve().parent / "polices")) + "/"
pdfmetrics.registerFont(TTFont("DV", FD + "DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DVB", FD + "DejaVuSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("DVI", FD + "DejaVuSans-Oblique.ttf"))
W, H = A4


# ------------------------------------------------------------------ helpers
class Doc:
    def __init__(self, name, title=None, author="Basira Conseil"):
        self.c = canvas.Canvas(str(OUT / name), pagesize=A4)
        self.c.setTitle(title or name)
        self.c.setAuthor(author)
        self.y = H - 25 * mm

    def text(self, t, size=10, font="DV", gap=1.45, x=20 * mm, color=(0.1, 0.1, 0.1)):
        self.c.setFont(font, size)
        self.c.setFillColorRGB(*color)
        for line in t.split("\n"):
            if self.y < 25 * mm:
                self.c.showPage()
                self.y = H - 25 * mm
                self.c.setFont(font, size)
            self.c.drawString(x, self.y, line)
            self.y -= size * gap

    def space(self, n=4):
        self.y -= n * mm

    def rule(self):
        self.c.setStrokeColorRGB(0.75, 0.75, 0.75)
        self.c.line(20 * mm, self.y, W - 20 * mm, self.y)
        self.y -= 5 * mm

    def table(self, rows, cols, size=9, bold_first=True):
        for i, r in enumerate(rows):
            self.c.setFont("DVB" if (i == 0 and bold_first) else "DV", size)
            for x, cell in zip(cols, r):
                if cell.startswith(">"):
                    self.c.drawRightString(x * mm, self.y, cell[1:])
                else:
                    self.c.drawString(x * mm, self.y, cell)
            self.y -= size * 1.7

    def footer(self, t=FICTIF):
        self.c.setFont("DVI", 7)
        self.c.setFillColorRGB(0.45, 0.45, 0.45)
        self.c.drawString(20 * mm, 12 * mm, t)

    def save(self):
        self.footer()
        self.c.save()


BASIRA = "Basira Conseil SAS · Paris 10e · contact@basira-conseil.example"


def invoice(fname, num, date, client_block, lines, ht, tva_label, tva, ttc, due, note=None, paid=None,
            issuer=("Basira Conseil SAS", "Paris 10e", "comptabilite@basira-conseil.example")):
    d = Doc(fname, f"Facture {num}")
    d.text(issuer[0], 15, "DVB")
    d.text(f"{issuer[1]} · {issuer[2]}", 8)
    d.space(6)
    d.text(f"FACTURE {num}", 18, "DVB")
    d.text(f"Date d'émission : {date}        Échéance : {due}", 9)
    d.space(4)
    d.text("Facturé à :", 9, "DVB")
    d.text(client_block, 9)
    d.space(6)
    d.table([["Désignation", "Qté", ">Montant HT"]] + [[a, b, ">" + c] for a, b, c in lines], [20, 130, 190])
    d.space(2)
    d.rule()
    d.table([["", "Total HT", ">" + ht], ["", tva_label, ">" + tva], ["", "Total TTC", ">" + ttc]], [20, 120, 190], bold_first=False)
    d.space(6)
    if note:
        d.text(note, 8, "DVI")
    if paid:
        d.space(4)
        d.text(paid, 11, "DVB", color=(0.1, 0.5, 0.2))
    d.space(8)
    d.text("Paiement par virement à 30 jours. Coordonnées bancaires communiquées séparément.", 8)
    d.save()


# ------------------------------------------------------------------ client briefs (duplicate pair)
BRIEF = f"""# Transports Ardelis — Brief atelier
_{FICTIF}_ · Notes de Yasmine, lundi 7 septembre 2026

**Client :** Transports Ardelis, transport routier de marchandises régional, Saint-Priest (Lyon), 240 salariés.
**Sponsor :** Sarah Martin, directrice des opérations.
**Informatique :** Julien Morel, responsable informatique.

## Besoin
Un atelier d'une journée, en présentiel, pour l'équipe de direction des opérations, sur l'usage de l'IA pour deux sujets : les devis de transport et les réponses aux emails clients.

## Format pressenti
- Présentiel, sur leur site de Saint-Priest.
- Date à caler, plutôt mi-octobre.
- « Une douzaine » de responsables d'exploitation. Liste à venir.

## Contraintes
- Aucune donnée client dans un outil externe pendant l'atelier.
- L'informatique doit valider tous les outils utilisés.

## Points ouverts
- Budget : pas encore abordé.
- Suite après l'atelier : pas abordée.
"""
(OUT / "Brief Ardelis.md").write_text(BRIEF, encoding="utf-8")
(OUT / "Brief Ardelis (1).md").write_text(BRIEF, encoding="utf-8")


# ------------------------------------------------------------------ docx
def docx_file(name, paragraphs):
    dd = docx.Document()
    st = dd.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for kind, text in paragraphs:
        if kind == "h1":
            dd.add_heading(text, 1)
        elif kind == "h2":
            dd.add_heading(text, 2)
        elif kind == "b":
            dd.add_paragraph(text, style="List Bullet")
        elif kind == "i":
            p = dd.add_paragraph()
            p.add_run(text).italic = True
        elif kind == "table":
            t = dd.add_table(rows=0, cols=len(text[0]))
            t.style = "Table Grid"
            for row in text:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = v
        else:
            dd.add_paragraph(text)
    dd.core_properties.author = "Basira Conseil"
    dd.save(str(OUT / name))


docx_file("Proposition Ardelis v1.docx", [
    ("i", f"BROUILLON — non envoyé · lundi 28 septembre 2026 · {FICTIF}"),
    ("h1", "Proposition — Atelier « IA pour les opérations »"),
    ("p", "Client : Transports Ardelis · À l'attention de Sarah Martin, directrice des opérations"),
    ("p", "Préparée par : Yasmine Haddad, Basira Conseil"),
    ("h2", "Périmètre"),
    ("b", "Atelier dirigeants d'une journée, jusqu'à 12 participants, sur site à Saint-Priest"),
    ("b", "Deux exercices sur mesure sur données anonymisées : devis de transport, réponses aux emails clients"),
    ("b", "Suivi post-atelier : deux demi-journées dans le mois suivant"),
    ("h2", "Budget"),
    ("table", [["Poste", "Montant HT"], ["Atelier dirigeants (1 jour, jusqu'à 12 participants)", "9 500 €"],
               ["Exercices sur mesure", "5 500 €"], ["Suivi post-atelier (2 demi-journées)", "3 000 €"],
               ["Total", "18 000 € HT"]]),
    ("p", "Date proposée : mercredi 14 ou jeudi 15 octobre 2026 (à confirmer)."),
    ("p", "Conditions : acompte de 50 % à la signature, solde à 30 jours après l'atelier."),
])

docx_file("programme atelier brouillon.docx", [
    ("i", f"Brouillon Inès — mardi 6 octobre 2026 — pas encore envoyé · {FICTIF}"),
    ("h1", "Atelier IA pour les opérations — Transports Ardelis — programme détaillé"),
    ("p", "Jeudi 15 octobre 2026 · Salle Rhône · Saint-Priest"),
    ("b", "09:00 — Accueil, objectifs de la journée (Sarah Martin)"),
    ("b", "09:30 — Ce que les assistants IA savent faire, et ne savent pas faire"),
    ("b", "10:30 — Exercice 1 : préparer un devis de transport sur données anonymisées"),
    ("b", "12:30 — Déjeuner sur place"),
    ("b", "13:30 — Exercice 2 : répondre à des emails clients"),
    ("b", "15:30 — Règles d'usage : uniquement les outils validés par l'informatique"),
    ("b", "16:30 — Plan d'action par équipe"),
    ("b", "17:15 — Clôture"),
    ("p", "Question ouverte : faut-il une deuxième salle pour le groupe service client pendant l'exercice 2 ?"),
])

docx_file("notes reunion equipe 05-10.docx", [
    ("i", f"Notes de réunion d'équipe — lundi 5 octobre 2026 — prises par Léa · {FICTIF}"),
    ("p", "Présents : Yasmine, Inès, Mehdi, Paul, Amina, Hugo, Léa."),
    ("h2", "Transports Ardelis"),
    ("b", "Proposition v2 signée vendredi (16,5 k€ HT). Atelier le jeudi 15 octobre à Saint-Priest."),
    ("b", "Sarah Martin veut le programme détaillé pour vendredi 9. Inès finalise le programme, Yasmine l'envoie."),
    ("b", "Amina prépare une première version des deux exercices de devis d'ici lundi 12."),
    ("b", "Visite du site demain mardi (Yasmine)."),
    ("h2", "Al Rawiya Holding"),
    ("b", "Appel de découverte mercredi 7 à 15 h au bureau de Paris (Yasmine + Mehdi)."),
    ("b", "Mehdi propose de caler un appel avec leur DAF le vendredi 16."),
    ("b", "Décision : aucun prix avant d'avoir parlé au décideur."),
    ("h2", "Maison Corvelle"),
    ("b", "Comité de pilotage mercredi 7 à 10 h (Paul + Yasmine)."),
    ("b", "Hugo : facture d'août toujours impayée, relance ferme à valider par Yasmine."),
    ("h2", "Divers"),
    ("b", "Léa : newsletter d'octobre à valider la semaine prochaine."),
    ("b", "Mehdi : un assistant IA pour trier les emails ? On en reparle au prochain point."),
])


# ------------------------------------------------------------------ xlsx
def xlsx(name, header, rows, widths, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title or "Feuille1"
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F3B52")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    ws.append([])
    ws.append([FICTIF])
    wb.properties.creator = "Basira Conseil"
    wb.save(str(OUT / name))


xlsx("participants Ardelis.xlsx", ["Prénom", "Nom", "Fonction", "Équipe"], [
    ["Julie", "Bernard", "Responsable d'exploitation", "Route régionale"],
    ["Thomas", "Petit", "Responsable d'exploitation", "Route régionale"],
    ["Inès", "Robert", "Responsable d'exploitation", "Entrepôt"],
    ["Lucas", "Richard", "Responsable d'exploitation", "Entrepôt"],
    ["Sophie", "Durand", "Responsable d'exploitation", "Longue distance"],
    ["Hugo", "Leroy", "Responsable d'exploitation", "Longue distance"],
    ["Camille", "Moreau", "Chargée de tarification", "Tarification"],
    ["Nathan", "Simon", "Chargé de tarification", "Tarification"],
    ["Léa", "Laurent", "Responsable d'exploitation", "International"],
    ["Louis", "Lefebvre", "Responsable d'exploitation", "International"],
    ["Chloé", "Michel", "Responsable planification", "Planification"],
    ["Arthur", "Garcia", "Responsable planification", "Planification"],
    ["Manon", "David", "Responsable service client", "Service client"],
    ["Jules", "Bertrand", "Conseiller service client", "Service client"],
], [12, 14, 30, 20], "Participants")

xlsx("Basira tarifs 2026.xlsx", ["Code", "Offre", "Contenu", "Prix HT (€)", "Unité"], [
    ["DIAG", "Diagnostic stratégique", "10 jours : entretiens, analyse, restitution au comité de direction", 12000, "forfait"],
    ["ATEL", "Atelier dirigeants (1 jour)", "Jusqu'à 12 participants, préparation et appel de préparation de 30 min inclus", 9500, "forfait"],
    ["ATEL+", "Participant supplémentaire", "Au-delà de 12, maximum 16", 500, "par personne"],
    ["EXO", "Exercices sur mesure", "2 exercices sur données du client anonymisées", 5500, "forfait"],
    ["FORM", "Formation équipe (2 jours)", "Jusqu'à 15 participants", 14000, "forfait"],
    ["COACH", "Accompagnement mensuel", "2 demi-journées + 1 comité de pilotage par mois, engagement 6 mois", 4800, "par mois"],
    ["DEPL", "Déplacements", "France métropolitaine incluse, hors France sur devis", 0, "—"],
    [], ["Conditions", "Acompte 50 % à la signature (missions ponctuelles), solde à 30 jours. Remise : accord écrit de Yasmine Haddad, 10 % maximum."],
], [8, 28, 62, 12, 14], "Grille 2026")


# ------------------------------------------------------------------ PDFs
# signed proposal v2
d = Doc("Proposition Ardelis v2 SIGNEE.pdf", "Proposition Transports Ardelis v2 signée")
d.text("Basira Conseil", 16, "DVB")
d.text(BASIRA, 8)
d.space(6)
d.text("PROPOSITION — Atelier « IA pour les opérations » — version 2", 13, "DVB")
d.text("Client : Transports Ardelis · À l'attention de Sarah Martin, directrice des opérations", 9)
d.text("Date : mercredi 30 septembre 2026 · Validité : 30 jours", 9)
d.space(4)
d.text("Périmètre", 11, "DVB")
d.text("• Atelier dirigeants d'une journée le jeudi 15 octobre 2026, sur site à Saint-Priest\n"
       "• Jusqu'à 15 participants\n"
       "• Deux exercices sur mesure sur données anonymisées : devis de transport, réponses aux emails clients\n"
       "• Appel de préparation de 30 minutes avec la sponsor et l'informatique", 9)
d.space(4)
d.text("Budget", 11, "DVB")
d.table([["Poste", ">Montant HT"], ["Atelier dirigeants (1 jour, jusqu'à 12 participants)", ">9 500 €"],
         ["3 participants supplémentaires (13 à 15) × 500 €", ">1 500 €"], ["Exercices sur mesure", ">5 500 €"]], [20, 190])
d.rule()
d.table([["Total", ">16 500 € HT"]], [20, 190])
d.space(3)
d.text("Conditions : acompte de 50 % (8 250 € HT) à la signature, solde à 30 jours après l'atelier.\n"
       "Déplacements en France métropolitaine inclus. Non inclus : accompagnement post-atelier, licences d'outils.", 9)
d.space(10)
d.text("Bon pour accord", 10, "DVB")
d.text("Sarah Martin, directrice des opérations, Transports Ardelis", 9)
d.text("Signé électroniquement le vendredi 2 octobre 2026", 9, "DVI")
d.save()

invoice("F-2026-026.pdf", "F-2026-026", "1er septembre 2026",
        "Dar Al Rimal\nÀ l'attention de Youssef Karam, directeur de cabinet du président\nDubaï, Émirats arabes unis",
        [("Diagnostic stratégique — solde (50 %)", "1", "6 000,00 €")], "6 000,00 €", "TVA", "0,00 €", "6 000,00 €",
        "1er octobre 2026", note="Exonération de TVA : prestation de services rendue à un preneur établi hors de l'Union européenne.",
        paid="PAYÉE le 22 septembre 2026")
invoice("FACTURE ACOMPTE ARDELIS.pdf", "F-2026-028", "2 octobre 2026",
        "Transports Ardelis\nService comptabilité\nSaint-Priest, France",
        [("Atelier dirigeants « IA pour les opérations » — acompte 50 % (proposition v2 du 30/09/2026)", "1", "8 250,00 €")],
        "8 250,00 €", "TVA 20 %", "1 650,00 €", "9 900,00 €", "à réception")
invoice("Corvelle F-2026-024.pdf", "F-2026-024", "3 août 2026",
        "Maison Corvelle\nÀ l'attention de Nadia Belkacem, responsable administrative et financière\nBeaune, France",
        [("Accompagnement mensuel — août 2026", "1", "4 800,00 €")], "4 800,00 €", "TVA 20 %", "960,00 €", "5 760,00 €",
        "2 septembre 2026")
invoice("facture 2024 old.pdf", "IL-2024-0187", "12 mars 2024", "Basira Conseil SAS\nParis 10e",
        [("Impression de 200 plaquettes A4 recto verso", "200", "340,00 €")], "340,00 €", "TVA 20 %", "68,00 €", "408,00 €",
        "11 avril 2024", paid="ACQUITTÉE", issuer=("Imprimerie Lorvanne (fictif)", "Montreuil", "contact@lorvanne.example"))

d = Doc("note de frais Lyon.pdf", "Note de frais")
d.text("Note de frais — Yasmine Haddad", 14, "DVB")
d.text("Basira Conseil · Mois : octobre 2026", 9)
d.space(4)
d.table([["Date", "Objet", "Client", ">Montant TTC"],
         ["06/10/2026", "Train Paris – Lyon Part-Dieu (aller)", "Transports Ardelis", ">89,00 €"],
         ["06/10/2026", "Train Lyon Part-Dieu – Paris (retour)", "Transports Ardelis", ">89,00 €"],
         ["06/10/2026", "Taxi gare – Saint-Priest", "Transports Ardelis", ">32,00 €"]], [20, 45, 125, 190])
d.rule()
d.table([["", "", "Total", ">210,00 €"]], [20, 45, 125, 190], bold_first=False)
d.space(4)
d.text("Objet : visite du site avant l'atelier du 15 octobre. Justificatifs joints.", 9)
d.save()

d = Doc("RIB Basira Conseil.pdf", "Relevé d'identité bancaire")
d.text("RELEVÉ D'IDENTITÉ BANCAIRE", 14, "DVB")
d.space(4)
d.text("Titulaire : BASIRA CONSEIL SAS", 10)
d.text("Banque : Banque Fictive de Démonstration", 10)
d.text("IBAN : FR76 0000 0000 0000 0000 0000 000 (FICTIF)", 10)
d.text("BIC : DEMOFRPPXXX (FICTIF)", 10)
d.space(6)
d.text("Document confidentiel. Ne pas transmettre.", 9, "DVI")
d.save()

d = Doc("Outils autorises IT.pdf", "Outils autorisés — atelier du 15/10", "Transports Ardelis — DSI")
d.text("Transports Ardelis — Direction informatique", 13, "DVB")
d.text("Outils d'IA autorisés pour l'atelier du jeudi 15 octobre 2026 · Émis par Julien Morel le 6 octobre 2026", 9)
d.space(4)
d.table([["Outil", "Statut", "Condition"],
         ["Microsoft Copilot (environnement entreprise)", "Autorisé", "Comptes Transports Ardelis uniquement"],
         ["ChatGPT Enterprise", "Autorisé pour l'atelier", "Données anonymisées uniquement"],
         ["Comptes personnels ChatGPT / Claude / Gemini", "Interdit", "—"],
         ["Extensions de navigateur", "Interdit", "—"]], [20, 100, 140])
d.space(4)
d.text("Comptes temporaires : créés le mercredi 14 octobre 2026.\nContact le jour J : Julien Morel.", 9)
d.save()

d = Doc("Organigramme Al Rawiya.pdf", "Organigramme division distribution", "Al Rawiya Holding")
d.text("Al Rawiya Holding — Organigramme simplifié", 13, "DVB")
d.text("Division distribution · document transmis le 7 octobre 2026", 9)
d.space(5)
for line in ["Président du groupe — Abdulaziz Al Rawiya",
             "  ├── Directeur financier groupe — Khalid Al-Rashid (Riyad) · valide les budgets",
             "  ├── Directrice de la transformation — Sarah Al-Mansouri (Paris)",
             "  └── Directeur de la division distribution — Fahad Al-Otaibi (Riyad)",
             "        ├── Région Arabie saoudite — 31 magasins",
             "        ├── Région Émirats — 11 magasins",
             "        └── Approvisionnement et logistique (centralisé à Riyad)"]:
    d.text(line, 10)
d.save()

d = Doc("Basira one-pager.pdf", "Basira Conseil — présentation")
d.text("Basira Conseil", 20, "DVB")
d.text("Voir clair avant d'agir.", 12, "DVI")
d.space(5)
d.text("Conseil en stratégie et en transformation pour les dirigeants d'ETI et de groupes familiaux,\nen France et dans les pays du Golfe.", 10)
d.space(5)
d.text("Ce que nous faisons", 11, "DVB")
d.text("• Diagnostic stratégique\n• Ateliers dirigeants\n• Formation des équipes\n• Accompagnement mensuel des dirigeants", 10)
d.space(4)
d.text("Notre approche", 11, "DVB")
d.text("Un interlocuteur senior du début à la fin. Des décisions mieux prises, pas des jours-hommes.", 10)
d.space(4)
d.text("contact@basira-conseil.example · Paris", 9)
d.save()

# scanned site-visit notes: image-only PDF
img = Image.new("L", (1240, 1754), 242)
dr = ImageDraw.Draw(img)
hand = ImageFont.truetype(FD + "DejaVuSans-Oblique.ttf", 34)
lines = ["Visite site Ardelis — 06/10", "", "Salle Rhône 2e étage : 16 places,", "1 vidéoprojecteur, pas de 2e écran",
         "", "WIFI invités : bloque ChatGPT / Copilot", "→ Julien doit ouvrir l'accès avant le 14/10", "",
         "Comptes temporaires : remis le 14/10 (à confirmer)", "", "Parking visiteurs : 6 places", "Déjeuner sur place 12h30",
         "", "2e salle service client ? → voir Sarah M.", "", "(fictif)"]
for i, l in enumerate(lines):
    dr.text((120 + random.Random(i).randint(-6, 6), 170 + i * 72), l, fill=45, font=hand)
img = img.rotate(-1.4, fillcolor=242).filter(ImageFilter.GaussianBlur(0.8))
noise = Image.effect_noise(img.size, 18).convert("L")
img = Image.blend(img, noise, 0.08)
img.convert("RGB").save(str(OUT / "scan_0013.pdf"), resolution=150)

# ------------------------------------------------------------------ noise
rnd = random.Random(7)
photo = Image.new("RGB", (1600, 1067))
pd = ImageDraw.Draw(photo)
for y in range(1067):
    t = y / 1067
    col = (int(70 + 120 * t), int(140 + 60 * t), int(200 - 40 * t)) if y < 600 else (int(200 - 60 * t), int(180 - 40 * t), int(120 - 20 * t))
    pd.line([(0, y), (1600, y)], fill=col)
pd.ellipse([1180, 120, 1320, 260], fill=(250, 220, 150))
photo = photo.filter(ImageFilter.GaussianBlur(2))
photo.save(str(OUT / "IMG_2041.jpg"), quality=82)

shot = Image.new("RGB", (1440, 900), (246, 247, 249))
sd = ImageDraw.Draw(shot)
f = ImageFont.truetype(FD + "DejaVuSans.ttf", 22)
fb = ImageFont.truetype(FD + "DejaVuSans-Bold.ttf", 26)
sd.rectangle([0, 0, 1440, 60], fill=(47, 59, 82))
sd.text((30, 16), "Tableau de bord — trafic du site (fictif)", fill="white", font=fb)
for i, h in enumerate([220, 260, 240, 310, 380, 350, 420]):
    sd.rectangle([120 + i * 160, 760 - h, 220 + i * 160, 760], fill=(88, 134, 201))
sd.text((120, 790), "Visites hebdomadaires — semaines 33 à 39", fill=(90, 90, 90), font=f)
shot.save(str(OUT / "Capture d’écran 2026-10-01 à 18.42.png"))

buf = io.BytesIO()
with zipfile.ZipFile(OUT / "Archive.zip", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("old/notes reunion 2023-11.txt", "Notes de réunion de novembre 2023 — sujets internes. (fictif)\n")
    z.writestr("old/ancien logo.txt", "Ancien logo Basira, version 2021 — fichier remplacé. (fictif)\n")
print("binary files done")
